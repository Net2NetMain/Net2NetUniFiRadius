import io
import json

from openpyxl import Workbook, load_workbook

from app import main
from test_api import client, sign_in  # shared isolated appliance fixture


def upload(client, content, filename='users.csv', kind='users', duplicates='skip', token=None):
    return client.post('/api/import/' + ('apply' if token else 'preview'),
        params={'filename': filename, 'kind': kind, 'duplicates': duplicates}, content=content,
        headers={'Content-Type': 'application/octet-stream', **({'X-Import-Preview': token} if token else {})})


def users_csv(password='Imported-Password-123!', username='alice', vlan=20):
    return f'username,display_name,password,profile,vlan,status\n{username},Alice,{password},Staff,{vlan},active\n'.encode()


def test_csv_preview_apply_backup_and_replay(client):
    assert upload(client, users_csv()).status_code == 401
    sign_in(client)
    data = users_csv()
    result = upload(client, data).json()
    assert result['valid'] and result['counts']['created'] == 2
    assert 'Imported-Password' not in json.dumps(result)
    assert client.get('/api/state').json()['profiles'] == []
    applied = upload(client, data, token=result['token']).json()
    assert applied['applied']
    assert (main.DATA / 'imports' / applied['backup']).exists()
    state = client.get('/api/state').json()
    assert len(state['users']) == len(state['profiles']) == 1
    assert upload(client, data, token=result['token']).status_code == 409


def test_bad_rows_are_atomic_and_no_secret_echo(client):
    sign_in(client)
    data = users_csv() + b'bob,Bob,short6,Staff,20,active\n'
    result = upload(client, data).json()
    assert not result['valid'] and result['errors'][0]['row'] == 3
    assert 'short6' not in json.dumps(result)
    result = upload(client, data, token=result['token']).json()
    assert not result.get('applied')
    assert client.get('/api/state').json()['users'] == []
    assert client.get('/api/state').json()['profiles'] == []


def test_duplicate_skip_update_and_blank_password(client):
    sign_in(client)
    data = users_csv()
    preview = upload(client, data).json()
    upload(client, data, token=preview['token'])
    with main.db() as c:
        old_password = c.execute('SELECT password FROM users').fetchone()[0]
    updated = b'username,display_name,password,profile,vlan,status\nalice,New name,,Staff,20,disabled\n'
    preview = upload(client, updated).json()
    assert preview['counts']['skipped'] == 1
    upload(client, updated, token=preview['token'])
    assert client.get('/api/state').json()['users'][0]['status'] == 'active'
    preview = upload(client, updated, duplicates='update').json()
    assert preview['counts']['updated'] == 1
    assert upload(client, updated, duplicates='update', token=preview['token']).json()['applied']
    assert client.get('/api/state').json()['users'][0]['status'] == 'disabled'
    with main.db() as c:
        assert c.execute('SELECT password FROM users').fetchone()[0] == old_password


def test_import_seven_character_password(client):
    sign_in(client)
    data = users_csv(password='7041839')
    preview = upload(client, data).json()
    assert preview['valid']
    assert upload(client, data, token=preview['token']).json()['applied']


def test_full_excel_and_templates(client):
    sign_in(client)
    workbook = Workbook()
    workbook.active.title = 'users'
    workbook.active.append(['username', 'password', 'profile'])
    workbook.active.append(['alice', 'Imported-Password-123!', 'Staff'])
    profile = workbook.create_sheet('profiles')
    profile.append(['name', 'vlan'])
    profile.append(['Staff', 20])
    devices = workbook.create_sheet('devices')
    devices.append(['name', 'address', 'secret'])
    devices.append(['Office AP', '192.0.2.10', 'Imported-Secret-123!'])
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    preview = upload(client, data, 'system.xlsx', 'system').json()
    assert preview['valid'], preview
    assert preview['counts']['created'] == 3
    assert upload(client, data, 'system.xlsx', 'system', token=preview['token']).json()['applied']
    state = client.get('/api/state').json()
    assert len(state['users']) == len(state['profiles']) == len(state['nas']) == 1
    template = client.get('/api/import/template/system?format=xlsx')
    assert load_workbook(io.BytesIO(template.content)).sheetnames == ['profiles', 'users', 'devices']


def test_import_rollback_and_stale_preview(client, monkeypatch):
    sign_in(client)
    data = users_csv()
    preview = upload(client, data).json()
    assert upload(client, users_csv(username='bob'), token=preview['token']).status_code == 409
    client.post('/api/profiles', json={'name': 'Guest', 'vlan': 30})
    assert upload(client, data, token=preview['token']).status_code == 409
    preview = upload(client, data).json()
    def fail(c):
        raise RuntimeError('invalid config')
    monkeypatch.setattr(main, 'apply_config', fail)
    assert upload(client, data, token=preview['token']).status_code == 503
    state = client.get('/api/state').json()
    assert state['users'] == []
    assert [p['name'] for p in state['profiles']] == ['Guest']


def test_reject_formulas_oversized_and_duplicate_rows(client):
    sign_in(client)
    data = users_csv() + users_csv().split(b'\n')[1] + b'\n'
    assert not upload(client, data).json()['valid']
    assert upload(client, b'not a workbook', 'system.xlsx', 'system').status_code == 422
    assert upload(client, b'a' * (8 * 1024 * 1024 + 1)).status_code == 413
    workbook = Workbook()
    workbook.active.title = 'profiles'
    workbook.active.append(['name', 'vlan'])
    workbook.active.append(['=1+1', 20])
    output = io.BytesIO()
    workbook.save(output)
    assert upload(client, output.getvalue(), 'system.xlsx', 'system').status_code == 422

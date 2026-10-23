"""Bounded CSV/XLSX imports with authenticated preview and transactional application."""
import csv
import hashlib
import hmac
import io
import json
import sqlite3
import subprocess
import time
import zipfile
from contextlib import closing

from fastapi import Depends, HTTPException, Request, Response
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from starlette.concurrency import run_in_threadpool

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 5000
HEADERS = {
    'profiles': ['name', 'vlan'],
    'users': ['username', 'display_name', 'password', 'profile', 'vlan', 'status'],
    'devices': ['name', 'address', 'secret'],
}
ALIASES = {'user_name': 'username', 'user': 'username', 'full_name': 'display_name',
           'profile_name': 'profile', 'vlan_id': 'vlan', 'shared_secret': 'secret',
           'ip_address': 'address', 'nas_ip_address': 'address', 'pass': 'password'}


class ImportFormatError(ValueError):
    pass


def normal_header(value):
    value = str(value or '').strip().lower().replace(' ', '_').replace('-', '_')
    return ALIASES.get(value, value)


def rows_from_table(kind, table):
    try:
        headers = [normal_header(v) for v in next(table)]
    except StopIteration:
        raise ImportFormatError(f'{kind}: missing header row')
    # Excel may retain formatting in trailing empty columns.
    while headers and not headers[-1]:
        headers.pop()
    if not headers or len(headers) > 30 or len(set(headers)) != len(headers):
        raise ImportFormatError(f'{kind}: headers must be present, unique and at most 30 columns')
    allowed = set(HEADERS[kind])
    unknown = set(headers) - allowed
    if unknown:
        raise ImportFormatError(f'{kind}: unrecognised columns; download the template for supported headers')
    required = {'profiles': {'name', 'vlan'}, 'users': {'username', 'profile'},
                'devices': {'name', 'address'}}[kind]
    if not required <= set(headers):
        raise ImportFormatError(f'{kind}: required headers: {", ".join(sorted(required))}')
    result = []
    for number, values in enumerate(table, 2):
        if number > MAX_ROWS + 2:
            raise ImportFormatError(f'Import limit is {MAX_ROWS} rows')
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if len(values) > len(headers) and any(str(v or '').strip() for v in values[len(headers):]):
            raise ImportFormatError(f'{kind} row {number}: more values than headers')
        row = {key: '' if value is None else str(value) for key, value in zip(headers, values)}
        for key in row:
            if key not in ('password', 'secret'):
                row[key] = row[key].strip()
        if any(len(v) > 512 for v in row.values()):
            raise ImportFormatError(f'{kind} row {number}: cell exceeds 512 characters')
        result.append((kind, number, row))
        if len(result) > MAX_ROWS:
            raise ImportFormatError(f'Import limit is {MAX_ROWS} rows')
    return result


def parse_file(data, filename, kind):
    if not data or len(data) > MAX_BYTES:
        raise ImportFormatError('Choose a non-empty file up to 8 MB')
    if filename.lower().endswith('.csv'):
        if kind not in HEADERS:
            raise ImportFormatError('Select users, profiles or devices for a CSV file')
        try:
            text = data.decode('utf-8-sig')
        except UnicodeError:
            raise ImportFormatError('Save the CSV as UTF-8, or use an .xlsx workbook')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return rows_from_table(kind, iter(csv.reader(io.StringIO(text), dialect, strict=True)))
    if not filename.lower().endswith('.xlsx'):
        raise ImportFormatError('Supported files: .csv and .xlsx. Save older .xls files as .xlsx first.')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 1000 or sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
            raise ImportFormatError('Workbook is too large when expanded')
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    result = []
    try:
        found = set()
        for sheet in workbook:
            sheet_kind = sheet.title.strip().lower()
            sheet_kind = {'nas': 'devices', 'network devices': 'devices', 'vlan profiles': 'profiles',
                          'wi-fi users': 'users'}.get(sheet_kind, sheet_kind)
            if sheet_kind not in HEADERS:
                continue
            if kind != 'system' and sheet_kind != kind:
                continue
            if sheet_kind in found:
                raise ImportFormatError('Use only one sheet for each record type')
            found.add(sheet_kind)
            if sheet.max_column and sheet.max_column > 30:
                raise ImportFormatError('Workbook sheets may contain at most 30 columns')
            def values(sheet=sheet):
                for cells in sheet.iter_rows():
                    if any(c.data_type == 'f' for c in cells):
                        raise ImportFormatError('Formula cells are not supported; paste values before importing')
                    yield [c.value for c in cells]
            result.extend(rows_from_table(sheet_kind, iter(values())))
        if not found:
            raise ImportFormatError('No matching sheets. Use users, profiles and/or devices sheet names.')
    finally:
        workbook.close()
    if len(result) > MAX_ROWS:
        raise ImportFormatError(f'Import limit is {MAX_ROWS} rows across all sheets')
    order = {'profiles': 0, 'users': 1, 'devices': 2}
    return sorted(result, key=lambda row: order[row[0]])


def snapshot(c):
    values = [list(map(tuple, c.execute(f'SELECT * FROM {table} ORDER BY id')))
              for table in ('profiles', 'users', 'nas')]
    return hashlib.sha256(json.dumps(values).encode()).hexdigest()


def preview_token(data, kind, duplicates, fingerprint):
    from . import main as m
    message = hashlib.sha256(data).hexdigest() + ':' + kind + ':' + duplicates + ':' + fingerprint
    return hmac.new((m.DATA / 'secret.key').read_bytes(), message.encode(), hashlib.sha256).hexdigest()


def process_rows(c, rows, duplicates):
    from . import main as m
    errors, changes, seen = [], [], set()
    counts = {'created': 0, 'updated': 0, 'skipped': 0}
    for kind, number, row in rows:
        try:
            identity = row.get('username' if kind == 'users' else 'name', '')
            if not identity:
                raise ValueError('Username/name is required')
            if (kind, identity) in seen:
                raise ValueError('Duplicate name/username within this file')
            seen.add((kind, identity))
            table = 'nas' if kind == 'devices' else kind
            key = 'username' if kind == 'users' else 'name'
            existing = c.execute(f'SELECT * FROM {table} WHERE {key}=?', (identity,)).fetchone()
            if existing and duplicates == 'skip':
                action = 'skipped'
            elif kind == 'profiles':
                model = m.Profile.model_validate(row)
                if existing:
                    c.execute('UPDATE profiles SET vlan=? WHERE id=?', (model.vlan, existing['id']))
                else:
                    c.execute('INSERT INTO profiles(name,vlan) VALUES (?,?)', (model.name, model.vlan))
                action = 'updated' if existing else 'created'
            elif kind == 'devices':
                model = m.Nas.model_validate(row)
                if not model.secret and not existing:
                    raise ValueError('New devices need a shared secret')
                secret = m.encryptor().encrypt(model.secret.encode()).decode() if model.secret else existing['secret']
                if existing:
                    c.execute('UPDATE nas SET address=?,secret=? WHERE id=?', (model.address, secret, existing['id']))
                else:
                    c.execute('INSERT INTO nas(name,address,secret) VALUES (?,?,?)', (model.name, model.address, secret))
                action = 'updated' if existing else 'created'
            else:
                profile_name = row.get('profile', '')
                profile = c.execute('SELECT * FROM profiles WHERE name=?', (profile_name,)).fetchone()
                vlan = row.get('vlan', '')
                if profile and vlan and int(vlan) != profile['vlan']:
                    raise ValueError('VLAN conflicts with the selected profile')
                if not profile:
                    model = m.Profile(name=profile_name, vlan=vlan)
                    c.execute('INSERT INTO profiles(name,vlan) VALUES (?,?)', (model.name, model.vlan))
                    profile = c.execute('SELECT * FROM profiles WHERE name=?', (profile_name,)).fetchone()
                    changes.append({'type': 'profiles', 'row': number, 'name': profile_name, 'action': 'created', 'automatic': True})
                    counts['created'] += 1
                model = m.User(username=identity, display_name=row.get('display_name') or identity,
                               password=row.get('password', ''), profile_id=profile['id'],
                               status=row.get('status') or 'active')
                if not model.password and not existing:
                    raise ValueError('New users need a password')
                password = m.encryptor().encrypt(model.password.encode()).decode() if model.password else existing['password']
                if existing:
                    c.execute('UPDATE users SET display_name=?,password=?,profile_id=?,status=? WHERE id=?',
                              (model.display_name, password, model.profile_id, model.status, existing['id']))
                else:
                    c.execute('INSERT INTO users(username,display_name,password,profile_id,status) VALUES (?,?,?,?,?)',
                              (identity, model.display_name, password, model.profile_id, model.status))
                action = 'updated' if existing else 'created'
            counts[action] += 1
            if action == 'skipped':
                details = 'Existing record kept'
            elif kind == 'profiles':
                details = f'VLAN {model.vlan}'
            elif kind == 'devices':
                details = model.address
            else:
                details = f'{profile_name} / VLAN {profile["vlan"]} / {model.status}'
            changes.append({'type': kind, 'row': number, 'name': identity, 'action': action, 'details': details})
        except ValidationError as exc:
            # Never include Pydantic's input values (passwords/shared secrets) in errors.
            message = '; '.join(f'{".".join(map(str, e["loc"]))}: {e["msg"]}' for e in exc.errors())
            errors.append({'type': kind, 'row': number, 'message': message})
        except sqlite3.IntegrityError:
            errors.append({'type': kind, 'row': number, 'message': 'Conflicting device IP or duplicate record'})
        except (ValueError, TypeError):
            errors.append({'type': kind, 'row': number, 'message':
                           'Check required fields, VLAN/profile consistency, password and shared secret'})
        if len(errors) >= 100:
            break
    return {'counts': counts, 'changes': changes, 'errors': errors, 'rows': len(rows),
            'valid': bool(rows) and not errors}


def execute_import(data, filename, kind, duplicates, token=None):
    from . import main as m
    if kind not in (*HEADERS, 'system') or duplicates not in ('skip', 'update'):
        raise HTTPException(422, 'Invalid import type or duplicate handling')
    try:
        rows = parse_file(data, filename, kind)
    except ImportFormatError as exc:
        raise HTTPException(422, str(exc))
    except Exception:
        # Parser exceptions can contain raw cells. Return a safe message for malformed files.
        raise HTTPException(422, 'Cannot read this file. Use the CSV/XLSX templates, plain values and supported headers; limit 8 MB / 5,000 rows.')
    with m.LOCK, m.db() as c:
        fingerprint = snapshot(c)
        expected = preview_token(data, kind, duplicates, fingerprint)
        if token is not None and not hmac.compare_digest(token, expected):
            raise HTTPException(409, 'File, options or system records changed. Preview the import again.')
        # Reserve a write transaction so no other writer changes the preview snapshot.
        c.execute('BEGIN IMMEDIATE')
        result = process_rows(c, rows, duplicates)
        c.rollback()
        result['token'] = expected
        if token is None or not result['valid']:
            return result
        # Snapshot the previous database before changing any records.
        backup_dir = m.DATA / 'imports'
        backup_dir.mkdir(mode=0o700, exist_ok=True)
        backup_name = 'before-import-' + str(time.time_ns()) + '.sqlite'
        with closing(sqlite3.connect(backup_dir / backup_name)) as dest:
            c.backup(dest)
        (backup_dir / backup_name).chmod(0o600)
        c.execute('BEGIN IMMEDIATE')
        if snapshot(c) != fingerprint:
            raise HTTPException(409, 'Records changed. Preview again.')
        result = process_rows(c, rows, duplicates)
        if not result['valid']:
            c.rollback()
            return result
        try:
            m.apply_config(c)
        except (RuntimeError, OSError, subprocess.SubprocessError):
            raise HTTPException(503, 'RADIUS configuration failed; the entire import was rolled back')
        result['backup'] = backup_name
        result['applied'] = True
        return result


def template(kind, excel=False):
    if kind not in HEADERS and not (kind == 'system' and excel):
        raise HTTPException(404)
    if excel:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_kind in (HEADERS if kind == 'system' else [kind]):
            sheet = workbook.create_sheet(sheet_kind)
            sheet.append(HEADERS[sheet_kind])
            sheet.freeze_panes = 'A2'
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 24
                if column[0].value != 'vlan':
                    sheet.column_dimensions[column[0].column_letter].number_format = '@'
        output = io.BytesIO()
        workbook.save(output)
        return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename=net2net-{kind}-template.xlsx'})
    output = io.StringIO()
    csv.writer(output).writerow(HEADERS[kind])
    return Response(output.getvalue(), media_type='text/csv',
                    headers={'Content-Disposition': f'attachment; filename=net2net-{kind}-template.csv'})


def register_imports(app, admin):
    @app.get('/api/import/template/{kind}', dependencies=[Depends(admin)])
    def download_template(kind: str, format: str = 'csv'):
        return template(kind, excel=format == 'xlsx')

    async def receive(request, apply):
        data = bytearray()
        async for chunk in request.stream():
            data.extend(chunk)
            if len(data) > MAX_BYTES:
                raise HTTPException(413, 'Import file must be at most 8 MB')
        token = request.headers.get('X-Import-Preview') if apply else None
        if apply and not token:
            raise HTTPException(409, 'Preview the file before importing')
        return await run_in_threadpool(execute_import, bytes(data), request.query_params.get('filename', ''),
            request.query_params.get('kind', 'users'), request.query_params.get('duplicates', 'skip'), token)

    @app.post('/api/import/preview', dependencies=[Depends(admin)])
    async def preview(request: Request):
        return await receive(request, False)

    @app.post('/api/import/apply', dependencies=[Depends(admin)])
    async def apply(request: Request):
        return await receive(request, True)
